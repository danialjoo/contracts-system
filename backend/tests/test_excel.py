import io

from openpyxl import Workbook, load_workbook

from app import excel
from tests.conftest import contract_payload


def build_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append([
        "ردیف", "محل نگهداری", "تفاهم نامه/\nقرارداد", "  طرف قرارداد  ",
        "لینک اسکن اصل قرارداد/ تفاهم نامه", "الحاقیه", "واحد مربوطه",
        "شماره قرارداد/\nشماره تفاهم نامه", "تاریخ شماره گذاری قرارداد",
        "تاریخ شروع قرارداد/ \nتاریخ شروع تفاهم نامه", "مدت قرارداد/\nمدت تفاهم نامه",
        "تاریخ پایان قرارداد/ \nتاریخ پایان تفاهم نامه", "فرمول پایان ",
        "روزهای باقیمانده از قرارداد/وضعیت", "مدت زمان سپری شده", "موضوع قرارداد",
        "مبلغ قرارداد/\nمبلغ تامین مالی/ ریال", "تضامین قرارداد", "عوامل طرف قرارداد",
        "ساختگاه/ مکان/ محل اجرا", "ظرفیت", "سایر موارد", "Column1", "روزهای باقی مانده",
    ])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SAMPLE = [1, "بایگانی ۴/۱۲", "قرارداد", "ساتبا", "", "", "فنی و مهندسی",
          "۱۴۰۳/ق-۲۱۷", "۱۴۰۳/۰۴/۱۰", "۱۴۰۳/۰۴/۱۵", "۳۶ ماه", "۱۴۰۶/۰۴/۱۵",
          "", "", "", "احداث نیروگاه خورشیدی", "۲۴٬۵۰۰٬۰۰۰٬۰۰۰٬۰۰۰",
          "ضمانت‌نامه ۱۰٪", "مدیر پروژه", "ابرکوه، یزد", "۱۰۰ مگاوات", "", "", ""]


def test_headers_of_the_real_company_file_are_mapped():
    records, errors = excel.read_rows(build_xlsx([SAMPLE]))
    assert errors == []
    rec = records[0]
    assert rec["contract_number"] == "۱۴۰۳/ق-۲۱۷"
    assert rec["party"] == "ساتبا"
    assert rec["amount_rial"] == 24_500_000_000_000
    assert rec["start_jalali"] == "1403/04/15"
    assert rec["kind"] == "contract"


def test_computed_columns_are_ignored_not_imported():
    records, _ = excel.read_rows(build_xlsx([SAMPLE]))
    assert "days_remaining" not in records[0]
    assert excel.field_for("فرمول پایان ") is None
    assert excel.field_for("مدت زمان سپری شده") is None
    assert excel.field_for("Column1") is None


def test_header_only_file_yields_no_records():
    records, errors = excel.read_rows(build_xlsx([]))
    assert records == [] and errors == []


def test_row_without_number_is_reported():
    bad = list(SAMPLE)
    bad[7] = ""
    _, errors = excel.read_rows(build_xlsx([bad]))
    assert any("شماره قرارداد خالی" in e for e in errors)


def test_import_dry_run_writes_nothing(tech):
    payload = build_xlsx([SAMPLE])
    resp = tech.post("/api/imports/xlsx",
                     files={"file": ("in.xlsx", payload, "application/vnd.ms-excel")},
                     data={"default_section": "technical", "dry_run": "true"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["ready"] == 1 and body["imported"] == 0
    assert tech.get("/api/contracts").json()["total"] == 0


def test_import_commit_creates_rows_and_skips_duplicates(tech):
    payload = build_xlsx([SAMPLE])
    first = tech.post("/api/imports/xlsx",
                      files={"file": ("in.xlsx", payload, "application/vnd.ms-excel")},
                      data={"default_section": "technical", "dry_run": "false"}).json()
    assert first["imported"] == 1
    assert tech.get("/api/contracts").json()["total"] == 1

    second = tech.post("/api/imports/xlsx",
                       files={"file": ("in.xlsx", payload, "application/vnd.ms-excel")},
                       data={"default_section": "technical", "dry_run": "false"}).json()
    assert second["imported"] == 0
    assert any("از قبل در سامانه" in e for e in second["errors"])


def test_export_round_trips_through_import(tech):
    tech.post("/api/contracts", json=contract_payload())
    exported = tech.get("/api/contracts/export")
    assert exported.status_code == 200
    assert exported.headers["content-disposition"].startswith("attachment")

    wb = load_workbook(io.BytesIO(exported.content))
    ws = wb.active
    assert ws.sheet_view.rightToLeft is True
    headers = [c.value for c in ws[1]]
    assert "موضوع قرارداد" in headers and "بخش" in headers
    assert ws.cell(row=2, column=headers.index("موضوع قرارداد") + 1).value == \
        "احداث نیروگاه خورشیدی ابرکوه"

    records, errors = excel.read_rows(exported.content)
    assert errors == [] and len(records) == 1


def test_each_header_maps_to_its_own_field_not_a_substring_neighbour():
    """«عوامل طرف قرارداد» نباید با «طرف قرارداد» اشتباه گرفته شود."""
    assert excel.field_for("عوامل طرف قرارداد") == "counterparty_people"
    assert excel.field_for("  طرف قرارداد  ") == "party"
    assert excel.field_for("شماره قرارداد/\nشماره تفاهم نامه") == "contract_number"
    assert excel.field_for("تفاهم نامه/\nقرارداد") == "kind"
    assert excel.field_for("مدت قرارداد/\nمدت تفاهم نامه") == "duration_text"
    assert excel.field_for("تاریخ شروع قرارداد/ \nتاریخ شروع تفاهم نامه") == "start"
    assert excel.field_for("تاریخ پایان قرارداد/ \nتاریخ پایان تفاهم نامه") == "end"
    assert excel.field_for("تاریخ شماره گذاری قرارداد") == "numbered_on"
    assert excel.field_for("ساختگاه/ مکان/ محل اجرا") == "site"
    assert excel.field_for("محل نگهداری") == "storage_location"


def test_mou_kind_is_actually_detected():
    row = list(SAMPLE)
    row[2] = "تفاهم نامه"
    records, _ = excel.read_rows(build_xlsx([row]))
    assert records[0]["kind"] == "mou"
